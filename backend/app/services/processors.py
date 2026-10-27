from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from sqlalchemy.orm import Session

from app.adapters.asr import get_asr_provider
from app.adapters.vision import get_vision_provider
from app.models import Material, ProcessingJob


def _pptx_text_outline(path: Path) -> tuple[str, dict]:
    try:
        from pptx import Presentation
    except ImportError:
        return "", {"parser": "unavailable"}

    presentation = Presentation(path)
    slides: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = (shape.text or "").strip()
            if text:
                texts.append(text)
        slides.append(f"## 第 {index} 页\n" + "\n".join(texts))
    return "\n\n".join(slides), {"parser": "python-pptx", "slide_count": len(slides)}


def process_prepare_job(session: Session, job: ProcessingJob) -> None:
    material = session.get(Material, job.material_id)
    if not material:
        job.status = "failed"
        job.message = "关联文件不存在"
        job.completed_at = datetime.now()
        session.commit()
        return

    job.stage = "inspect"
    job.progress = 20
    job.message = "正在识别文件类型"
    material.status = "processing"
    session.commit()

    path = Path(material.stored_path)
    metadata = json.loads(material.metadata_json or "{}")

    if material.material_type == "presentation" and path.suffix.lower() == ".pptx":
        job.stage = "extract_ppt_text"
        job.progress = 55
        job.message = "正在提取 PPT 原生文字"
        session.commit()
        text, extraction_meta = _pptx_text_outline(path)
        material.extracted_text = text or None
        metadata.update(extraction_meta)
        material.status = "awaiting_recognition"
        job.status = "completed"
        job.stage = "recognition_pending"
        job.progress = 100
        job.message = "基础解析完成，等待配置 ASR 或视觉模型"
        job.completed_at = datetime.now()
    elif material.material_type == "audio":
        asr_provider = get_asr_provider()
        metadata.update({"recognition_provider": asr_provider.name})
        # 为音频创建后续的转录任务
        material.status = "awaiting_recognition"
        session.add(
            ProcessingJob(
                material_id=material.id,
                job_type="audio.transcribe",
                message=f"等待语音转写（使用 {asr_provider.name}）",
            )
        )
        job.status = "completed"
        job.stage = "recognition_pending"
        job.progress = 100
        job.message = "基础解析完成，已派发语音转写任务"
        job.completed_at = datetime.now()
    else:
        material.status = "ready"
        job.status = "completed"
        job.stage = "completed"
        job.progress = 100
        job.message = "基础元数据解析完成"
        job.completed_at = datetime.now()

    material.metadata_json = json.dumps(metadata, ensure_ascii=False)
    session.commit()


def process_transcribe_job(session: Session, job: ProcessingJob) -> None:
    material = session.get(Material, job.material_id)
    if not material:
        job.status = "failed"
        job.message = "关联文件不存在"
        job.completed_at = datetime.now()
        session.commit()
        return

    job.stage = "transcribing"
    job.progress = 30
    job.message = "正在执行语音转写"
    material.status = "processing"
    session.commit()

    path = Path(material.stored_path)
    if not path.is_file():
        job.status = "failed"
        job.message = "音频物理文件不存在"
        job.completed_at = datetime.now()
        session.commit()
        return

    asr = get_asr_provider()

    def on_progress(current: int, total: int, partial_text: str) -> None:
        # 切片进度映射到 30% ~ 95%
        percent = min(95, 30 + int((current / max(1, total)) * 65))
        job.progress = percent
        job.message = f"正在执行语音转写 ({current}/{total} 片)"
        try:
            session.commit()
        except Exception:
            session.rollback()

    try:
        result = asr.transcribe(path, progress_callback=on_progress)
        markdown_parts = ["# 语音转录文本", ""]
        if result.segments:
            for seg in result.segments:
                start_min, start_sec = divmod(seg.start_ms // 1000, 60)
                end_min, end_sec = divmod(seg.end_ms // 1000, 60)
                speaker_tag = f"【{seg.speaker}】" if seg.speaker else ""
                markdown_parts.append(f"**[{start_min:02d}:{start_sec:02d} - {end_min:02d}:{end_sec:02d}]** {speaker_tag}{seg.text}")
        else:
            markdown_parts.append(result.full_text)

        material.extracted_text = "\n\n".join(markdown_parts)
        metadata = json.loads(material.metadata_json or "{}")
        metadata["asr"] = {
            "provider": asr.name,
            "duration_ms": result.duration_ms,
            "language": result.language,
            "segment_count": len(result.segments),
        }
        material.metadata_json = json.dumps(metadata, ensure_ascii=False)
        material.status = "ready"

        job.status = "completed"
        job.stage = "completed"
        job.progress = 100
        job.message = f"转写完成（共 {len(result.segments)} 段）"
        job.completed_at = datetime.now()
        session.commit()
    except Exception as exc:
        material.status = "awaiting_recognition"
        job.status = "failed"
        job.message = f"转写失败：{str(exc)[:300]}"
        job.completed_at = datetime.now()
        session.commit()


def process_vision_job(session: Session, job: ProcessingJob) -> None:
    material = session.get(Material, job.material_id)
    if not material:
        job.status = "failed"
        job.message = "关联文件不存在"
        job.completed_at = datetime.now()
        session.commit()
        return

    job.stage = "analyzing"
    job.progress = 40
    job.message = "正在进行视觉模型分析"
    material.status = "processing"
    session.commit()

    vision = get_vision_provider()
    metadata = json.loads(material.metadata_json or "{}")
    metadata["vision_provider"] = vision.name
    material.metadata_json = json.dumps(metadata, ensure_ascii=False)
    material.status = "ready"

    job.status = "completed"
    job.stage = "completed"
    job.progress = 100
    job.message = "视觉理解分析完成"
    job.completed_at = datetime.now()
    session.commit()
