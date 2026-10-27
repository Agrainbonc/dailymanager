from __future__ import annotations

import asyncio
import os
import tempfile
import unittest
import zipfile
from pathlib import Path


class IngestionPipelineTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls._temp_dir = tempfile.TemporaryDirectory()
        os.environ["DAILY_MANAGER_DATA_DIR"] = str(Path(cls._temp_dir.name) / "app-data")

        from app.config import ensure_data_directories
        from app.database import Base, SessionLocal, engine
        import app.models  # Register SQLAlchemy mappings before creating the test database.

        ensure_data_directories()
        cls.Base = Base
        cls.SessionLocal = SessionLocal
        cls.engine = engine
        Base.metadata.create_all(bind=engine)

    @classmethod
    def tearDownClass(cls) -> None:
        cls._temp_dir.cleanup()

    def setUp(self) -> None:
        from app.models import AppSetting, Course, LessonSession, Material, ProcessingJob

        with self.SessionLocal() as session:
            session.query(ProcessingJob).delete()
            session.query(Material).delete()
            session.query(LessonSession).delete()
            session.query(Course).delete()
            session.query(AppSetting).delete()
            session.commit()

    def test_scan_archives_presentation_and_runs_basic_parse(self) -> None:
        from pptx import Presentation

        from app.models import Course, Material, ProcessingJob
        from app.services.job_runner import job_runner
        from app.services.library import scan_course_watch_path

        drop_folder = Path(self._temp_dir.name) / "drop"
        drop_folder.mkdir()
        pptx_path = drop_folder / "高等数学_2026-09-16_第一讲.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "极限"
        slide.placeholders[1].text = "函数极限的定义"
        deck.save(pptx_path)

        with self.SessionLocal() as session:
            course = Course(name="高等数学", watch_path=str(drop_folder))
            session.add(course)
            session.commit()

            imported = scan_course_watch_path(session, course)
            session.commit()
            self.assertEqual(len(imported), 1)
            self.assertTrue(Path(imported[0].stored_path).is_file())

        self.assertTrue(job_runner._process_next_job())

        with self.SessionLocal() as session:
            material = session.query(Material).one()
            job = session.query(ProcessingJob).one()
            self.assertEqual(material.status, "awaiting_recognition")
            self.assertIn("极限", material.extracted_text or "")
            self.assertEqual(job.status, "completed")
            self.assertEqual(job.stage, "recognition_pending")

    def test_inbox_keyword_classification_moves_file_into_course_folder(self) -> None:
        from app.models import Course
        from app.services.library import scan_input_directory, update_library_settings

        inbox = Path(self._temp_dir.name) / "keyword-inbox"
        inbox.mkdir()
        source = inbox / "卷积_2026年09月16日.m4a"
        source.write_bytes(b"keyword-classification")

        with self.SessionLocal() as session:
            course = Course(name="信号与系统", keywords_json='["卷积"]')
            session.add(course)
            session.flush()
            update_library_settings(session, input_directory=str(inbox), output_directory=None)
            imported = scan_input_directory(session)
            session.commit()

            self.assertEqual(len(imported), 1)
            material = imported[0]
            self.assertEqual(material.course_id, course.id)
            self.assertEqual(material.classification_status, "classified")
            self.assertGreaterEqual(material.classification_confidence, 66)
            self.assertFalse(source.exists())
            self.assertTrue(Path(material.stored_path).is_file())
            self.assertEqual(Path(material.stored_path).parents[2].name, "信号与系统")

    def test_inbox_uses_filename_time_and_schedule_when_name_has_no_course(self) -> None:
        from app.models import Course
        from app.services.library import scan_input_directory, update_library_settings

        inbox = Path(self._temp_dir.name) / "schedule-inbox"
        inbox.mkdir()
        source = inbox / "课堂录音_2026-09-14_09-10.wav"
        source.write_bytes(b"schedule-classification")

        with self.SessionLocal() as session:
            course = Course(
                name="数据结构",
                schedule_json='[{"weekday": 1, "start_time": "08:00", "end_time": "10:00"}]',
            )
            session.add(course)
            session.flush()
            update_library_settings(session, input_directory=str(inbox), output_directory=None)
            imported = scan_input_directory(session)
            session.commit()

            self.assertEqual(len(imported), 1)
            self.assertEqual(imported[0].course_id, course.id)
            self.assertEqual(imported[0].classification_status, "classified")
            self.assertIn("课表匹配", imported[0].classification_reason)

    def test_ambiguous_schedule_requires_manual_assignment_and_moves_file(self) -> None:
        from app.models import Course
        from app.services.library import assign_material_to_course, scan_input_directory, update_library_settings

        inbox = Path(self._temp_dir.name) / "review-inbox"
        inbox.mkdir()
        source = inbox / "课堂录音_2026-09-14_09-00.wav"
        source.write_bytes(b"ambiguous-classification")

        with self.SessionLocal() as session:
            first_course = Course(
                name="算法设计",
                schedule_json='[{"weekday": 1, "start_time": "08:00", "end_time": "10:00"}]',
            )
            second_course = Course(
                name="计算机网络",
                schedule_json='[{"weekday": 1, "start_time": "08:00", "end_time": "10:00"}]',
            )
            session.add_all([first_course, second_course])
            session.flush()
            update_library_settings(session, input_directory=str(inbox), output_directory=None)
            imported = scan_input_directory(session)
            session.commit()

            self.assertEqual(len(imported), 1)
            material = imported[0]
            self.assertIsNone(material.course_id)
            self.assertEqual(material.classification_status, "needs_review")
            self.assertIn("需要确认", material.classification_reason)
            self.assertIn("needs-review", Path(material.stored_path).parts)

            assign_material_to_course(session, material, first_course)
            session.commit()

            self.assertEqual(material.course_id, first_course.id)
            self.assertEqual(material.classification_status, "manual_assigned")
            self.assertNotIn("needs-review", Path(material.stored_path).parts)
            self.assertTrue(Path(material.stored_path).is_file())

    def test_course_update_replaces_classification_rules_without_changing_legacy_path(self) -> None:
        from pydantic import ValidationError

        from app.api import update_course
        from app.models import Course
        from app.schemas import CourseUpdate

        with self.SessionLocal() as session:
            course = Course(
                name="工程力学基础",
                aliases_json='["工程力学"]',
                schedule_json='[{"weekday": 2, "start_time": "14:20", "end_time": "16:00"}]',
                watch_path="D:/legacy/engineering-mechanics",
            )
            session.add(course)
            session.commit()

            updated = update_course(
                course.id,
                CourseUpdate(
                    aliases=["工程力学", "工程力学", "Mechanics"],
                    keywords=["静力学", "静力学"],
                    schedule=[
                        {"weekday": 2, "start_time": "14:20", "end_time": "16:00"},
                        {"weekday": 5, "start_time": "16:30", "end_time": "18:10"},
                    ],
                ),
                session,
            )

            self.assertEqual(updated.aliases, ["工程力学", "Mechanics"])
            self.assertEqual(updated.keywords, ["静力学"])
            self.assertEqual(len(updated.schedule), 2)
            self.assertEqual(course.watch_path, "D:/legacy/engineering-mechanics")

        with self.assertRaises(ValidationError):
            CourseUpdate(schedule=[{"weekday": 1, "start_time": "10:00", "end_time": "08:00"}])

    def test_course_export_contains_source_and_text_packages(self) -> None:
        from app.models import Course, Material
        from app.services.exports import create_course_export, remove_export

        material_path = Path(self._temp_dir.name) / "slides.pptx"
        material_path.write_bytes(b"presentation-source")
        with self.SessionLocal() as session:
            course = Course(name="英语学术交流")
            session.add(course)
            session.flush()
            parsed_material = Material(
                course_id=course.id,
                material_type="presentation",
                status="awaiting_recognition",
                classification_status="classified",
                classification_confidence=90,
                classification_reason="课程名“英语学术交流”",
                original_filename="Academic English.pptx",
                source_path=str(material_path),
                stored_path=str(material_path),
                sha256="a" * 64,
                file_size=material_path.stat().st_size,
                inferred_by="test",
                extracted_text="## 第 1 页\nAcademic presentation skills",
            )
            pending_material = Material(
                course_id=course.id,
                material_type="audio",
                status="awaiting_recognition",
                classification_status="classified",
                classification_confidence=68,
                classification_reason="课表匹配",
                original_filename="课堂录音.m4a",
                source_path=str(material_path),
                stored_path=str(Path(self._temp_dir.name) / "missing.m4a"),
                sha256="b" * 64,
                file_size=10,
                inferred_by="test",
            )
            session.add_all([parsed_material, pending_material])
            session.commit()

            source_export = create_course_export(session, course, "source")
            text_export = create_course_export(session, course, "text")

        try:
            with zipfile.ZipFile(source_export) as archive:
                self.assertIn("manifest.json", archive.namelist())
                source_entry = next(name for name in archive.namelist() if name.startswith("source/"))
                self.assertEqual(archive.read(source_entry), b"presentation-source")

            with zipfile.ZipFile(text_export) as archive:
                self.assertIn("README.md", archive.namelist())
                texts = [name for name in archive.namelist() if name.startswith("texts/")]
                self.assertEqual(len(texts), 2)
                exported_text = "\n".join(archive.read(name).decode("utf-8") for name in texts)
                self.assertIn("Academic presentation skills", exported_text)
                self.assertIn("尚未产生可导出的识别文本", exported_text)
        finally:
            remove_export(source_export)
            remove_export(text_export)

    def test_inbox_clean_and_moves_to_default_library_root_when_no_output_dir(self) -> None:
        from app.config import LIBRARY_DIR
        from app.models import Course
        from app.services.library import scan_input_directory, update_library_settings

        inbox = Path(self._temp_dir.name) / "clean-inbox"
        inbox.mkdir()
        sub_folder = inbox / "nested-box"
        sub_folder.mkdir()
        source_file = sub_folder / "高数_2026-09-14.wav"
        source_file.write_bytes(b"clean-inbox-test")

        with self.SessionLocal() as session:
            course = Course(name="高等数学", aliases_json='["高数"]')
            session.add(course)
            session.flush()
            update_library_settings(session, input_directory=str(inbox), output_directory=None)
            imported = scan_input_directory(session)
            session.commit()

            self.assertEqual(len(imported), 1)
            material = imported[0]
            # 存储位置应在系统的 LIBRARY_DIR 下，而非仍在投递箱中
            self.assertTrue(Path(material.stored_path).is_relative_to(LIBRARY_DIR))
            # 投递箱原文件已被搬移，且空子目录已被清理
            self.assertFalse(source_file.exists())
            self.assertFalse(sub_folder.exists())

    def test_custom_course_export_filters_and_aligns_text(self) -> None:
        from app.models import Course, Material
        from app.services.exports import create_custom_course_export, remove_export

        material_path = Path(self._temp_dir.name) / "lecture_recording.m4a"
        material_path.write_bytes(b"recording-data")

        with self.SessionLocal() as session:
            course = Course(name="操作系统原理")
            session.add(course)
            session.flush()

            m1 = Material(
                course_id=course.id,
                material_type="audio",
                status="ready",
                classification_status="classified",
                classification_confidence=95,
                classification_reason="课程名匹配",
                original_filename="第1讲.m4a",
                source_path=str(material_path),
                stored_path=str(material_path),
                sha256="c" * 64,
                file_size=material_path.stat().st_size,
                inferred_by="test",
                extracted_text="### 进程与线程的区别\n1. 进程是资源分配单位",
            )
            m2 = Material(
                course_id=course.id,
                material_type="audio",
                status="ready",
                classification_status="classified",
                classification_confidence=90,
                classification_reason="课程名匹配",
                original_filename="第2讲.m4a",
                source_path=str(material_path),
                stored_path=str(material_path),
                sha256="d" * 64,
                file_size=material_path.stat().st_size,
                inferred_by="test",
                extracted_text="### 内存虚拟化",
            )
            session.add_all([m1, m2])
            session.commit()

            # 仅勾选 m1，并同时勾选源文件与 ASR 文本
            archive_path = create_custom_course_export(
                session=session,
                course=course,
                material_ids=[m1.id],
                include_source=True,
                include_text=True,
            )

        try:
            with zipfile.ZipFile(archive_path) as archive:
                namelist = archive.namelist()
                self.assertIn("manifest.json", namelist)
                self.assertIn("README.md", namelist)
                # 只有 m1 被导出
                sources = [n for n in namelist if n.startswith("source/")]
                texts = [n for n in namelist if n.startswith("texts/")]
                self.assertEqual(len(sources), 1)
                self.assertEqual(len(texts), 1)
                self.assertTrue(sources[0].endswith("第1讲.m4a"))
                content = archive.read(texts[0]).decode("utf-8")
                self.assertIn("进程是资源分配单位", content)
        finally:
            remove_export(archive_path)

    def test_get_material_file_endpoint(self) -> None:
        from fastapi import HTTPException
        from fastapi.responses import FileResponse
        from app.api import get_material_file
        from app.models import Course, Material

        audio_file = Path(self._temp_dir.name) / "test_sample.mp3"
        audio_file.write_bytes(b"ID3FakeMp3Content")

        with self.SessionLocal() as session:
            course = Course(name="操作系统")
            session.add(course)
            session.commit()

            material = Material(
                course_id=course.id,
                material_type="audio",
                status="ready",
                classification_status="classified",
                classification_confidence=95,
                classification_reason="测试",
                original_filename="sample.mp3",
                source_path=str(audio_file),
                stored_path=str(audio_file),
                sha256="e" * 64,
                file_size=audio_file.stat().st_size,
                inferred_by="test",
            )
            session.add(material)
            session.commit()
            m_id = material.id

            # 测试正常获取/试听
            res = asyncio.run(get_material_file(material_id=m_id, download=False, session=session))
            self.assertIsInstance(res, FileResponse)
            self.assertEqual(res.media_type, "audio/mpeg")

            # 测试下载
            res_dl = asyncio.run(get_material_file(material_id=m_id, download=True, session=session))
            self.assertIn("attachment", res_dl.headers.get("content-disposition", ""))

            # 测试不存在的资料
            with self.assertRaises(HTTPException) as ctx:
                asyncio.run(get_material_file(material_id=99999, session=session))
            self.assertEqual(ctx.exception.status_code, 404)

    def test_job_runner_stale_recovery_and_retry(self) -> None:
        from app.models import Course, Material, ProcessingJob
        from app.services.job_runner import job_runner

        with self.SessionLocal() as session:
            course = Course(name="队列测试课程")
            session.add(course)
            session.commit()

            material = Material(
                course_id=course.id,
                material_type="audio",
                status="processing",
                original_filename="stale.mp3",
                source_path="dummy.mp3",
                stored_path="dummy.mp3",
                sha256="f" * 64,
                file_size=1024,
            )
            session.add(material)
            session.commit()

            # 创建模拟服务中断遗留的 running 任务
            stale_job = ProcessingJob(
                material_id=material.id,
                job_type="audio.transcribe",
                status="running",
                stage="transcribing",
                progress=30,
                message="正在执行语音转写",
                attempts=1,
            )
            session.add(stale_job)
            session.commit()
            stale_job_id = stale_job.id

        # 运行恢复机制
        recovered = job_runner.recover_stale_jobs(requeue=True)
        self.assertGreaterEqual(recovered, 1)

        # 校验恢复后状态
        with self.SessionLocal() as session:
            job = session.get(ProcessingJob, stale_job_id)
            mat = session.get(Material, material.id)
            self.assertIsNotNone(job)
            self.assertEqual(job.status, "queued")
            self.assertEqual(job.progress, 0)
            self.assertIn("服务重启恢复排队", job.message)
            self.assertEqual(mat.status, "queued")

            # 模拟任务失败后再手动重试
            job.status = "failed"
            job.message = "转写网络超时"
            mat.status = "awaiting_recognition"
            session.commit()

        # 手动重试
        retry_ok = job_runner.retry_job(stale_job_id)
        self.assertTrue(retry_ok)

        with self.SessionLocal() as session:
            retried_job = session.get(ProcessingJob, stale_job_id)
            mat = session.get(Material, material.id)
            self.assertEqual(retried_job.status, "queued")
            self.assertEqual(mat.status, "queued")

    def test_asr_progress_callback(self) -> None:
        from app.adapters.asr import MockAsrProvider

        provider = MockAsrProvider()
        calls = []

        def cb(current: int, total: int, partial: str) -> None:
            calls.append((current, total, partial))

        res = provider.transcribe(Path("test_audio.mp3"), progress_callback=cb)
        self.assertTrue(len(res.full_text) > 0)
        self.assertEqual(len(calls), 1)
        self.assertEqual(calls[0][0], 1)
        self.assertEqual(calls[0][1], 1)

    def test_list_materials_with_search_query(self) -> None:
        from app.api import list_materials
        from app.models import Course, Material

        with self.SessionLocal() as session:
            course = Course(name="数学分析")
            session.add(course)
            session.commit()

            m1 = Material(
                course_id=course.id,
                material_type="presentation",
                status="ready",
                original_filename="第一讲_极限论.pptx",
                source_path="dummy1.pptx",
                stored_path="dummy1.pptx",
                sha256="1" * 64,
                file_size=1024,
                extracted_text="关于柯西收敛准则的严格证明与定理分析。",
            )
            m2 = Material(
                course_id=course.id,
                material_type="audio",
                status="ready",
                original_filename="第2讲_导数.mp3",
                source_path="dummy2.mp3",
                stored_path="dummy2.mp3",
                sha256="2" * 64,
                file_size=2048,
                extracted_text="这一节讨论微分中值定理与拉格朗日公式。",
            )
            session.add_all([m1, m2])
            session.commit()

            # 按文件名搜
            results_name = list_materials(q="极限论", session=session)
            self.assertEqual(len(results_name), 1)
            self.assertEqual(results_name[0].original_filename, "第一讲_极限论.pptx")

            # 按正文文本搜
            results_text = list_materials(q="拉格朗日", session=session)
            self.assertEqual(len(results_text), 1)
            self.assertEqual(results_text[0].original_filename, "第2讲_导数.mp3")

            # 未命中
            results_none = list_materials(q="不存在的关键词xyz", session=session)
            self.assertEqual(len(results_none), 0)

    def test_upload_material_and_conflict_check(self) -> None:
        from app.models import Course, Material
        from app.services.library import save_uploaded_material

        fake_audio = b"ID3\x03\x00\x00\x00\x00\x00\x00fake audio content for test"

        with self.SessionLocal() as session:
            course = Course(name="操作系统原理")
            session.add(course)
            session.commit()

            # 1. 首次上传：无冲突，正常入库
            is_conflict, mat1, existing1 = save_uploaded_material(
                session=session,
                file_bytes=fake_audio,
                filename="lecture1.mp3",
                course=course,
                force_keep=False,
            )
            session.commit()
            self.assertFalse(is_conflict)
            self.assertIsNotNone(mat1)
            self.assertIsNone(existing1)
            mat1_id = mat1.id

            # 2. 再次上传相同文件（force_keep=False）：检测到冲突，返回已存在对象
            is_conflict2, mat2, existing2 = save_uploaded_material(
                session=session,
                file_bytes=fake_audio,
                filename="lecture1_copy.mp3",
                course=course,
                force_keep=False,
            )
            self.assertTrue(is_conflict2)
            self.assertIsNotNone(existing2)
            self.assertEqual(existing2.id, mat1_id)

            # 3. 用户确认保留副本（force_keep=True）：正常存为独立新副本
            is_conflict3, mat3, existing3 = save_uploaded_material(
                session=session,
                file_bytes=fake_audio,
                filename="lecture1_dup.mp3",
                course=course,
                force_keep=True,
            )
            session.commit()
            self.assertFalse(is_conflict3)
            self.assertIsNotNone(mat3)
            self.assertNotEqual(mat3.id, mat1_id)
            self.assertEqual(mat3.sha256, mat1.sha256)

    def test_duplicate_detection_and_deduplication(self) -> None:
        from app.models import Course, Material
        from app.services.library import find_duplicate_materials, deduplicate_materials

        with self.SessionLocal() as session:
            course = Course(name="线性代数")
            session.add(course)
            session.commit()

            # 创建具有相同 SHA256 的两份资料
            shared_sha = "d" * 64
            m_orig = Material(
                course_id=course.id,
                material_type="audio",
                status="ready",
                original_filename="线代_01.mp3",
                source_path="linear1.mp3",
                stored_path="linear1.mp3",
                sha256=shared_sha,
                file_size=5000,
                extracted_text="这是第一讲特征值与特征向量的完整转写",
            )
            m_dup = Material(
                course_id=course.id,
                material_type="audio",
                status="queued",
                original_filename="线代_01 (1).mp3",
                source_path="linear1_copy.mp3",
                stored_path="linear1_copy.mp3",
                sha256=shared_sha,
                file_size=5000,
                extracted_text=None,
            )
            session.add_all([m_orig, m_dup])
            session.commit()

            # 检测重复
            dups = find_duplicate_materials(session)
            found_group = [g for g in dups if g["sha256"] == shared_sha]
            self.assertEqual(len(found_group), 1)
            self.assertEqual(found_group[0]["count"], 2)

            # 执行去重清理
            cleaned, retained, freed = deduplicate_materials(session)
            self.assertGreaterEqual(cleaned, 1)
            self.assertGreaterEqual(retained, 1)

            # 验证去重后只保留有提取文本且 ready 的主项
            remaining = session.query(Material).filter(Material.sha256 == shared_sha).all()
            self.assertEqual(len(remaining), 1)
            self.assertEqual(remaining[0].id, m_orig.id)
            self.assertEqual(remaining[0].original_filename, "线代_01.mp3")

    def test_delete_material_with_ref_count_safety(self) -> None:
        from app.models import Course, Material, ProcessingJob
        from app.services.library import delete_material

        dummy_file = Path(self._temp_dir.name) / "test_shared_file.mp3"
        dummy_file.write_bytes(b"shared-audio-content")

        with self.SessionLocal() as session:
            course = Course(name="计算机网络")
            session.add(course)
            session.commit()

            # 创建两个 Material 指向同一个物理文件（共享存储引用）
            m1 = Material(
                course_id=course.id,
                material_type="audio",
                status="ready",
                original_filename="计网_01.mp3",
                source_path=str(dummy_file),
                stored_path=str(dummy_file),
                sha256="1" * 64,
                file_size=len(b"shared-audio-content"),
            )
            m2 = Material(
                course_id=course.id,
                material_type="audio",
                status="ready",
                original_filename="计网_01_副本.mp3",
                source_path=str(dummy_file),
                stored_path=str(dummy_file),
                sha256="1" * 64,
                file_size=len(b"shared-audio-content"),
            )
            session.add_all([m1, m2])
            session.commit()

            # 为 m1 挂载一个 ProcessingJob
            job = ProcessingJob(
                material_id=m1.id,
                job_type="transcribe",
                status="completed",
                stage="completed",
                progress=100,
            )
            session.add(job)
            session.commit()

            m1_id = m1.id
            m2_id = m2.id

            # 1. 删除 m1（物理删除开启）：因为 m2 仍在引用 dummy_file，所以物理文件不能被删除！
            res1 = delete_material(session, m1_id, delete_file=True)
            self.assertEqual(res1["material_id"], m1_id)
            self.assertFalse(res1["file_deleted"])  # 受到引用计数保护
            self.assertTrue(dummy_file.is_file())    # 磁盘文件必须完好无损

            # 关联的 job 也应被级联清理
            self.assertIsNone(session.get(ProcessingJob, job.id))
            self.assertIsNone(session.get(Material, m1_id))

            # 2. 删除 m2（最后一份引用）：此时物理文件应当被安全删除
            res2 = delete_material(session, m2_id, delete_file=True)
            self.assertEqual(res2["material_id"], m2_id)
            self.assertTrue(res2["file_deleted"])
            self.assertFalse(dummy_file.is_file())  # 文件已被删除
            self.assertIsNone(session.get(Material, m2_id))

    def test_delete_materials_batch(self) -> None:
        from app.models import Course, Material
        from app.services.library import delete_materials_batch

        f1 = Path(self._temp_dir.name) / "batch_1.txt"
        f2 = Path(self._temp_dir.name) / "batch_2.txt"
        f1.write_bytes(b"batch-content-1")
        f2.write_bytes(b"batch-content-2")

        with self.SessionLocal() as session:
            course = Course(name="操作系统")
            session.add(course)
            session.commit()

            m1 = Material(
                course_id=course.id,
                material_type="other",
                status="ready",
                original_filename="batch_1.txt",
                source_path=str(f1),
                stored_path=str(f1),
                sha256="2" * 64,
                file_size=15,
            )
            m2 = Material(
                course_id=course.id,
                material_type="other",
                status="ready",
                original_filename="batch_2.txt",
                source_path=str(f2),
                stored_path=str(f2),
                sha256="3" * 64,
                file_size=15,
            )
            session.add_all([m1, m2])
            session.commit()

            res = delete_materials_batch(session, [m1.id, m2.id], delete_files=True)
            self.assertEqual(res["deleted_count"], 2)
            self.assertEqual(set(res["deleted_material_ids"]), {m1.id, m2.id})
            self.assertFalse(f1.exists())
            self.assertFalse(f2.exists())
            self.assertIsNone(session.get(Material, m1.id))
            self.assertIsNone(session.get(Material, m2.id))

    def test_resolve_duplicate_group_specified_keeper(self) -> None:
        from app.models import Course, Material
        from app.services.library import resolve_duplicate_group

        with self.SessionLocal() as session:
            course = Course(name="离散数学")
            session.add(course)
            session.commit()

            sha = "e" * 64
            m_a = Material(
                course_id=course.id,
                material_type="audio",
                status="queued",
                original_filename="离散_A.mp3",
                source_path="a.mp3",
                stored_path="a.mp3",
                sha256=sha,
                file_size=3000,
            )
            m_b = Material(
                course_id=course.id,
                material_type="audio",
                status="ready",
                original_filename="离散_B.mp3",
                source_path="b.mp3",
                stored_path=str(Path(self._temp_dir.name) / "b.mp3"),
                sha256=sha,
                file_size=3000,
            )
            session.add_all([m_a, m_b])
            session.commit()

            # 明确指定 m_b 作为 keeper
            res = resolve_duplicate_group(session, sha256=sha, keeper_id=m_b.id)
            self.assertEqual(res["keeper_id"], m_b.id)
            self.assertEqual(res["removed_ids"], [m_a.id])

            remaining = session.query(Material).filter(Material.sha256 == sha).all()
            self.assertEqual(len(remaining), 1)
            self.assertEqual(remaining[0].id, m_b.id)

    def test_storage_audit_and_clean_orphan_files(self) -> None:
        from app.config import get_library_dir
        from app.models import Course, Material
        from app.services.library import audit_storage, clean_orphan_files

        # 在资料库目录下生成一个合法的和一个孤儿文件
        lib_dir = get_library_dir()
        lib_dir.mkdir(parents=True, exist_ok=True)
        valid_file = lib_dir / "valid_file.bin"
        orphan_file = lib_dir / "orphan_garbage.bin"
        valid_file.write_bytes(b"valid-data")
        orphan_file.write_bytes(b"orphan-data-garbage")

        with self.SessionLocal() as session:
            course = Course(name="编译原理")
            session.add(course)
            session.commit()

            mat = Material(
                course_id=course.id,
                material_type="other",
                status="ready",
                original_filename="valid_file.bin",
                source_path=str(valid_file),
                stored_path=str(valid_file),
                sha256="4" * 64,
                file_size=len(b"valid-data"),
            )
            session.add(mat)
            session.commit()

            # 1. 运行审计
            audit = audit_storage(session)
            self.assertGreaterEqual(audit["valid_files_count"], 1)
            self.assertGreaterEqual(audit["orphan_files_count"], 1)
            self.assertGreaterEqual(audit["orphan_files_bytes"], len(b"orphan-data-garbage"))

            # 2. 清理孤儿文件
            clean_res = clean_orphan_files(session)
            self.assertGreaterEqual(clean_res["cleaned_orphans_count"], 1)
            self.assertFalse(orphan_file.exists())  # 孤儿垃圾文件应被清除
            self.assertTrue(valid_file.exists())   # 资料库正常引用的文件必须保留！


if __name__ == "__main__":
    unittest.main()


